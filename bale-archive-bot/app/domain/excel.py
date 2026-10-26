"""Excel export: right-to-left sheets, Persian headers, Shamsi dates, clickable file links.

Sheet 1 «آرشیو»: one row per post; «فایل ۱»، «فایل ۲» … columns hold real
hyperlinks (build_media_url) titled «مشاهده‌ی فایل N».
Sheet 2 «فایل‌ها»: one row per file with size, duration, status and link.
Only SELECTs feed this module.
"""

from __future__ import annotations

import io
from datetime import datetime

import jdatetime
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet
from sqlalchemy.ext.asyncio import AsyncConnection

from app.config import Settings
from app.db.repo import MediaRow, media_for_posts
from app.domain.reports import ExportRow, PostFilter, export_rows
from app.i18n import fa
from app.mapping import POST_CONTENT_LABELS
from app.media_links import build_media_url, parse_bale_ref

HEADERS = [
    "شماره ثبت",
    "تاریخ ثبت (شمسی)",
    "تاریخ ارسال (شمسی)",
    "فرستنده",
    "شناسه بله",
    "گروه",
    "نوع محتوا",
    "هشتگ‌ها",
    "تعداد فایل",
    "متن",
]
WIDTHS = [10, 18, 18, 26, 14, 22, 12, 30, 10, 60]
FILE_HEADERS = [
    "شماره پست",
    "تاریخ (شمسی)",
    "فرستنده",
    "گروه",
    "هشتگ‌ها",
    "نوع",
    "حجم (مگابایت)",
    "مدت (ثانیه)",
    "وضعیت",
    "لینک مشاهده",
]
FILE_WIDTHS = [10, 18, 26, 22, 30, 10, 14, 12, 18, 30]
ON_SERVER = "روی سرور"
WAITING = "منتظر بارگذاری"


def shamsi(value: datetime | None) -> str:
    if value is None:
        return ""
    return jdatetime.datetime.fromgregorian(datetime=value).strftime("%Y/%m/%d %H:%M")


def view_label(index: int) -> str:
    return f"مشاهده‌ی فایل {fa.fa_digits(index)}"


def _header(sheet: Worksheet, headers: list[str], widths: list[int]) -> None:
    sheet.sheet_view.rightToLeft = True
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=1, column=index).column_letter].width = width
    sheet.freeze_panes = "A2"


def _link(cell: object, url: str, label: str) -> None:
    cell.value = label  # type: ignore[attr-defined]
    cell.hyperlink = url  # type: ignore[attr-defined]
    cell.style = "Hyperlink"  # type: ignore[attr-defined]


def build_workbook(
    rows: list[ExportRow],
    media: dict[int, list[MediaRow]] | None = None,
    settings: Settings | None = None,
    title: str = "آرشیو",
) -> bytes:
    media = media or {}
    max_files = max((len(media.get(row.post_id, [])) for row in rows), default=0)
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = title[:31]
    file_headers = [f"فایل {fa.fa_digits(i)}" for i in range(1, max_files + 1)]
    _header(sheet, HEADERS + file_headers, WIDTHS + [18] * max_files)
    for row in rows:
        sheet.append(
            [
                row.post_id,
                shamsi(row.created_at),
                shamsi(row.posted_at),
                row.sender,
                row.bale_user_id,
                row.group_name,
                POST_CONTENT_LABELS.get(row.content_type, str(row.content_type)),
                row.hashtags,
                row.files,
                row.content_text,
            ]
        )
        for index, item in enumerate(media.get(row.post_id, []), start=1):
            cell = sheet.cell(row=sheet.max_row, column=len(HEADERS) + index)
            _link(cell, build_media_url(item.id, settings), view_label(index))

    files = workbook.create_sheet("فایل‌ها")
    _header(files, FILE_HEADERS, FILE_WIDTHS)
    for row in rows:
        for item in media.get(row.post_id, []):
            ref = parse_bale_ref(item.storage_path)
            files.append(
                [
                    row.post_id,
                    shamsi(row.created_at),
                    row.sender,
                    row.group_name,
                    row.hashtags,
                    fa.media_kind(item.media_type),
                    round(item.file_size / 1024 / 1024, 2) if item.file_size else None,
                    item.duration,
                    WAITING if ref else ON_SERVER,
                    None,
                ]
            )
            cell = files.cell(row=files.max_row, column=len(FILE_HEADERS))
            if ref:
                cell.value = f"در گروه آرشیو بله – پیام {ref[1]}"
            else:
                _link(cell, build_media_url(item.id, settings), "مشاهده‌ی فایل")

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


async def build_export(
    conn: AsyncConnection, flt: PostFilter, settings: Settings | None = None
) -> tuple[bytes, int]:
    """(xlsx bytes, number of posts) for this filter."""
    rows = await export_rows(conn, flt)
    media = await media_for_posts(conn, [row.post_id for row in rows])
    return build_workbook(rows, media, settings), len(rows)
