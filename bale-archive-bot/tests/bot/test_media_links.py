"""Signed view links, Range streaming, manual upload of big files, Excel links, dry run."""

from __future__ import annotations

import io
import shutil
import subprocess
import sys
from pathlib import Path
from urllib.parse import unquote

import httpx
import pytest
from openpyxl import load_workbook

from app.domain.excel import build_export
from app.domain.reports import PostFilter
from app.media_links import build_media_url, build_upload_url, sign
from app.media_server import create_media_app
from tests.bot.harness import ADMIN_ID, ARCHIVE_ID, GROUP, MB, USER, Harness


def client(h: Harness) -> httpx.AsyncClient:
    transport = httpx.ASGITransport(app=create_media_app(h.ctx))
    return httpx.AsyncClient(transport=transport, base_url="http://95.38.191.89:8090")


async def _save_document(h: Harness, name: str, data: bytes, mime: str) -> int:
    h.register_group()
    h.fake.file_bytes["d1"] = data
    await h.send(GROUP, USER, document={"file_id": "d1", "file_name": name, "mime_type": mime, "file_size": len(data)})
    await h.confirm([1])
    await h.settle()
    (media_id,), = h.root.rows("SELECT id FROM PostMedia")
    return int(media_id)


async def test_view_link_returns_the_file(h: Harness) -> None:
    media_id = await _save_document(h, "گزارش.pdf", b"%PDF-1.4 hello", "application/pdf")
    url = build_media_url(media_id, h.settings)
    assert url.startswith("http://95.38.191.89:8090/m/")
    async with client(h) as c:
        response = await c.get(url)
    assert response.status_code == 200
    assert response.content == b"%PDF-1.4 hello"
    assert response.headers["content-type"] == "application/pdf"
    assert response.headers["content-disposition"].startswith("inline")
    assert response.headers["accept-ranges"] == "bytes"


async def test_wrong_signature_is_404(h: Harness) -> None:
    media_id = await _save_document(h, "a.pdf", b"%PDF-1.4", "application/pdf")
    async with client(h) as c:
        bad = await c.get(f"/m/{media_id}/{'0' * 32}")
        other = await c.get(f"/m/{media_id + 1}/{sign('view', media_id, 'test-link-secret')}")
        missing = await c.get(build_media_url(99999, h.settings))
    assert bad.status_code == other.status_code == missing.status_code == 404
    assert "پیدا نشد" in bad.text


async def test_range_request_returns_206(h: Harness) -> None:
    data = bytes(range(256)) * 4
    media_id = await _save_document(h, "a.pdf", data, "application/pdf")
    url = build_media_url(media_id, h.settings)
    async with client(h) as c:
        part = await c.get(url, headers={"Range": "bytes=10-19"})
        tail = await c.get(url, headers={"Range": "bytes=-5"})
        bad = await c.get(url, headers={"Range": "bytes=5000-"})
        head = await c.head(url)
    assert part.status_code == 206 and part.content == data[10:20]
    assert part.headers["content-range"] == f"bytes 10-19/{len(data)}"
    assert tail.status_code == 206 and tail.content == data[-5:]
    assert bad.status_code == 416
    assert head.status_code == 200 and head.headers["content-length"] == str(len(data))


async def test_docx_is_attachment_with_persian_name(h: Harness) -> None:
    docx = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    media_id = await _save_document(h, "صورت جلسه.docx", b"PK\x03\x04docx", docx)
    async with client(h) as c:
        response = await c.get(build_media_url(media_id, h.settings))
    disposition = response.headers["content-disposition"]
    assert response.headers["content-type"] == docx
    assert disposition.startswith("attachment")
    assert unquote(disposition.split("filename*=UTF-8''")[1]) == "صورت جلسه.docx"
    assert 'filename="post' in disposition


async def test_path_outside_media_folder_is_404(h: Harness, tmp_path: Path) -> None:
    media_id = await _save_document(h, "a.pdf", b"%PDF", "application/pdf")
    secret_file = tmp_path / "outside.txt"
    secret_file.write_text("secret")
    for path in (str(secret_file), str(h.settings.media_root / ".." / "outside.txt"), "/etc/passwd"):
        h.root.execute("UPDATE PostMedia SET storage_path = %s WHERE id = %s", (path, media_id))
        async with client(h) as c:
            response = await c.get(build_media_url(media_id, h.settings))
        assert response.status_code == 404


async def test_rate_limit_60_per_minute(h: Harness) -> None:
    media_id = await _save_document(h, "a.pdf", b"%PDF", "application/pdf")
    url = build_media_url(media_id, h.settings)
    async with client(h) as c:
        codes = [(await c.head(url)).status_code for _ in range(61)]
    assert codes[:60] == [200] * 60 and codes[60] == 429


async def _save_big_video(h: Harness) -> int:
    h.register_group()
    await h.send(GROUP, USER, video={"file_id": "big", "file_size": 30 * MB}, caption="ویدیوی بزرگ")
    await h.confirm([1])
    await h.settle()
    (media_id, path), = h.root.rows("SELECT id, storage_path FROM PostMedia")
    assert path.startswith(f"bale:{ARCHIVE_ID}/")
    return int(media_id)


async def test_big_file_view_page_and_admin_gets_upload_link(h: Harness) -> None:
    media_id = await _save_big_video(h)
    async with client(h) as c:
        page = await c.get(build_media_url(media_id, h.settings))
    assert page.status_code == 404 and "بزرگ‌تر از ۲۰ مگابایت" in page.text
    notices = [t for t in h.texts_to(ADMIN_ID) if t.startswith("📦")]
    assert len(notices) == 1 and build_upload_url(media_id, h.settings) in notices[0]
    # the panel lists it too
    await h.press({"id": ADMIN_ID, "is_bot": False, "first_name": "م"}, "pa", "pu")
    assert build_upload_url(media_id, h.settings) in h.texts_to(ADMIN_ID)[-1]


def _small_video(tmp_path: Path) -> bytes:
    out = tmp_path / "up.mp4"
    subprocess.run(
        ["ffmpeg", "-y", "-hide_banner", "-loglevel", "error", "-f", "lavfi", "-i",
         "testsrc=size=320x240:rate=10", "-t", "1", "-c:v", "libx264", "-pix_fmt", "yuv420p", str(out)],
        check=True,
    )  # fmt: skip
    return out.read_bytes()


@pytest.mark.skipif(shutil.which("ffprobe") is None, reason="ffprobe not installed")
async def test_upload_replaces_bale_reference(h: Harness, tmp_path: Path) -> None:
    from tests.conftest import EXECUTED_SQL

    media_id = await _save_big_video(h)
    data = _small_video(tmp_path)
    url = build_upload_url(media_id, h.settings)
    start = len(EXECUTED_SQL)
    async with client(h) as c:
        form = await c.get(url)
        assert form.status_code == 200 and "بارگذاری" in form.text
        done = await c.post(url, content=data, headers={"X-File-Name": "film.mp4"})
        assert done.status_code == 200 and "ذخیره شد" in done.text
        assert "۵٪" in done.text  # size differs a lot from the 30 MB Bale reported: warning
        await h.settle()
        (path, size, mime, width, height, duration), = h.root.rows(
            "SELECT storage_path, file_size, mime_type, width, height, duration FROM PostMedia"
        )
        assert path.startswith(str(h.settings.media_root)) and path.endswith(".mp4")
        assert mime == "video/mp4" and (width, height) == (320, 240) and duration == 1
        updates = [s for s in EXECUTED_SQL[start:] if s.lstrip().upper().startswith("UPDATE")]
        assert updates and all(s.startswith("UPDATE PostMedia SET") for s in updates)
        # the upload link is now dead, the view link serves the file
        again = await c.get(url)
        assert again.status_code == 404
        view = await c.get(build_media_url(media_id, h.settings))
        assert view.status_code == 200 and view.content == Path(path).read_bytes()


async def test_upload_wrong_signature_is_404(h: Harness) -> None:
    media_id = await _save_big_video(h)
    async with client(h) as c:
        response = await c.post(f"/u/{media_id}/{'f' * 32}", content=b"x")
        view_sig = build_media_url(media_id, h.settings).rsplit("/", 1)[1]
        swapped = await c.post(f"/u/{media_id}/{view_sig}", content=b"x")  # a view sig is no upload sig
    assert response.status_code == swapped.status_code == 404
    assert h.root.rows("SELECT storage_path FROM PostMedia")[0][0].startswith("bale:")


async def test_upload_over_60_mb_is_refused(h: Harness) -> None:
    media_id = await _save_big_video(h)

    async def body():  # type: ignore[no-untyped-def]
        chunk = b"\0" * MB
        for _ in range(61):
            yield chunk

    async with client(h) as c:
        response = await c.post(build_upload_url(media_id, h.settings), content=body())
    assert response.status_code == 413
    assert h.root.rows("SELECT storage_path FROM PostMedia")[0][0].startswith("bale:")
    assert list(h.settings.media_root.glob("upload-*")) == []


async def test_upload_of_wrong_type_is_refused(h: Harness) -> None:
    media_id = await _save_big_video(h)
    async with client(h) as c:
        response = await c.post(build_upload_url(media_id, h.settings), content=b"not a video at all")
    assert response.status_code == 415
    assert h.root.rows("SELECT storage_path FROM PostMedia")[0][0].startswith("bale:")


async def test_excel_links_are_build_media_url_and_serve_the_file(h: Harness) -> None:
    h.register_group()
    h.fake.file_bytes["p1"] = b"first-file"
    h.fake.file_bytes["p2"] = b"second-file"
    from app.bale.models import Update

    for message_id, file_id in ((8001, "p1"), (8002, "p2")):
        await h.dispatcher.dispatch(
            Update.model_validate(
                {"update_id": message_id, "message": {"message_id": message_id, "date": 1, "chat": GROUP,
                 "from": USER, "photo": [{"file_id": file_id, "file_size": 10}]}}
            )
        )  # fmt: skip
    await h.settle()
    await h.confirm([1])
    await h.settle()
    async with h.db.tx() as conn:
        data, count = await build_export(conn, PostFilter(), h.settings)
    assert count == 1
    book = load_workbook(io.BytesIO(data))
    main, files = book["آرشیو"], book["فایل‌ها"]
    assert main.sheet_view.rightToLeft and files.sheet_view.rightToLeft
    headers = [cell.value for cell in main[1]]
    assert headers[-2:] == ["فایل ۱", "فایل ۲"]
    ids = [int(r[0]) for r in h.root.rows("SELECT id FROM PostMedia ORDER BY id")]
    links = [main.cell(row=2, column=len(headers) - 1 + i) for i in range(2)]
    assert [c.value for c in links] == ["مشاهده‌ی فایل ۱", "مشاهده‌ی فایل ۲"]
    assert [c.hyperlink.target for c in links] == [build_media_url(i, h.settings) for i in ids]
    assert all(c.style == "Hyperlink" for c in links)
    assert [cell.value for cell in files[1]][-2:] == ["وضعیت", "لینک مشاهده"]
    assert files.cell(row=2, column=9).value == "روی سرور"
    async with client(h) as c:
        for cell, expected in zip(links, (b"first-file", b"second-file"), strict=True):
            response = await c.get(cell.hyperlink.target)
            assert response.status_code == 200 and response.content == expected


async def test_excel_waiting_file_text(h: Harness) -> None:
    await _save_big_video(h)
    async with h.db.tx() as conn:
        data, _ = await build_export(conn, PostFilter(), h.settings)
    files = load_workbook(io.BytesIO(data))["فایل‌ها"]
    row = [cell.value for cell in files[2]]
    assert row[8] == "منتظر بارگذاری"
    assert row[9].startswith("در گروه آرشیو بله – پیام ")


async def test_compress_existing_dry_run_changes_nothing(h: Harness) -> None:
    import importlib.util

    spec = importlib.util.spec_from_file_location(
        "compress_existing", Path(__file__).resolve().parents[2] / "scripts" / "compress_existing.py"
    )
    module = importlib.util.module_from_spec(spec)  # type: ignore[arg-type]
    sys.modules["compress_existing"] = module
    spec.loader.exec_module(module)  # type: ignore[union-attr]
    h.register_group()
    h.fake.file_bytes["v9"] = b"x" * 1000
    await h.send(GROUP, USER, video={"file_id": "v9", "file_size": 1000})
    await h.confirm([1])
    await h.settle()
    log_file = h.settings.data_path / "compression.json"
    log_file.unlink(missing_ok=True)  # pretend it was stored before compression existed
    before_rows = h.root.rows("SELECT id, storage_path, file_size, mime_type FROM PostMedia")
    before_files = {p: p.read_bytes() for p in h.settings.media_root.rglob("*") if p.is_file()}
    code = await module.run(apply=False, settings=h.settings)
    assert code == 0
    assert h.root.rows("SELECT id, storage_path, file_size, mime_type FROM PostMedia") == before_rows
    assert {p: p.read_bytes() for p in h.settings.media_root.rglob("*") if p.is_file()} == before_files
    assert not log_file.exists()
    if shutil.which("ffmpeg") and shutil.which("ffprobe"):
        assert await module.run(apply=True, settings=h.settings) == 0
        entry = next(iter(__import__("json").loads(log_file.read_text()).values()))
        assert entry["reason"] == "probe_failed"  # fake bytes are not a real video: kept as is
        assert h.root.rows("SELECT id, storage_path, file_size, mime_type FROM PostMedia") == before_rows
