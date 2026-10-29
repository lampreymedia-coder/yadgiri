"""CLAUDE-MERGE-SCHEMA.md §4: merged tables, shared Users rules, Source='bale', «📝 توضیحات»."""

from __future__ import annotations

import io
import re
from datetime import datetime

from openpyxl import load_workbook

from app.domain.excel import build_export
from app.domain.reports import PostFilter, build_report
from app.handlers.wizard import expire_and_remind
from app.i18n import fa
from app.media_links import build_media_url
from tests.bot.harness import GROUP, GROUP_ID, OTHER_USER, USER, Harness
from tests.conftest import EXECUTED_SQL

USER_COLUMNS = (
    "Username, PasswordHash, FirstName, LastName, Phone, BaleChatId, BaleUsername, IsAdmin, "
    "IsActor, MustChangePassword, BaleRequireCode, IsDeleted, IsActive, NotifyBale, "
    "NotifyMeetings, NotifyWeekly, NotifyDelays, FailedLogins, CreatedAt IS NOT NULL"
)


def _user_updates(since: int) -> list[str]:
    return [s for s in EXECUTED_SQL[since:] if re.match(r"\s*UPDATE\s+Users\b", s)]


async def _note_prompt(h: Harness, sid: str) -> int:
    await h.press(USER, "nt", sid)
    prompt = h.ctx.wizards.get(sid).note_prompt_id  # type: ignore[union-attr]
    assert prompt is not None
    assert fa.NOTE_PROMPT in h.texts_to(USER["id"])
    return prompt


async def _reply(h: Harness, prompt_id: int, body: str, sender: dict = USER, chat_id: int | None = None) -> int:  # type: ignore[type-arg]
    chat = {"id": chat_id or USER["id"], "type": "private"}
    return await h.send(chat, sender, text=body, reply_to_message={"message_id": prompt_id, "chat": chat, "date": 1})


# ─── Full save into the merged tables ───


async def test_new_bale_member_full_save(h: Harness) -> None:
    group_id = h.register_group(org_unit=7)
    h.fake.file_bytes["ph"] = b"photo-bytes"
    await h.send(GROUP, USER, photo=[{"file_id": "ph", "file_size": 11}], caption="کپشن")
    sid = h.sid()
    await h.press(USER, "y", sid)
    await h.press(USER, "t", sid, "2")
    prompt = await _note_prompt(h, sid)
    await _reply(h, prompt, "این عکس مربوط به جلسه‌ی دیروز است")
    await h.press(USER, "c", sid)
    await h.press(USER, "f", sid)
    await h.settle()

    assert h.root.rows(f"SELECT {USER_COLUMNS} FROM Users") == [
        (f"bale_{USER['id']}", "!bale-no-login", "علی", "رضایی", None, USER["id"], "ali",
         0, 0, 0, 0, 0, 1, 1, 0, 0, 0, 0, 1)
    ]  # fmt: skip
    (user_id,), = h.root.rows("SELECT Id FROM Users")
    posts = h.root.rows(
        "SELECT Id, AuthorUserId, BaleGroupId, OrgUnitId, Source, ContentType, Body, Title, UserNote "
        "FROM Posts"
    )
    assert len(posts) == 1
    post_id, author, bale_group, unit, source, ctype, body, title, note = posts[0]
    assert (author, bale_group, unit, source, ctype) == (user_id, group_id, 7, "bale", 2)
    assert (body, title, note) == ("کپشن", None, "این عکس مربوط به جلسه‌ی دیروز است")
    assert h.root.rows("SELECT PostId, HashtagId FROM PostHashtags") == [(post_id, 2)]
    assert h.root.rows("SELECT PostId, MediaType, BaleFileId FROM PostMedia") == [(post_id, 1, "ph")]
    assert h.root.rows("SELECT UserId, BaleGroupId, IsActive FROM BaleGroupMembers") == [
        (user_id, group_id, 1)
    ]


async def test_empty_first_and_last_name_defaults(h: Harness) -> None:
    h.register_group()
    nameless = {"id": 555, "is_bot": False, "first_name": ""}
    await h.text(GROUP, nameless, "سلام")
    assert h.root.rows("SELECT FirstName, LastName, Username, BaleUsername FROM Users") == [
        ("کاربر بله", "", "bale_555", None)
    ]


# ─── Users shared with the staff system ───


async def test_staff_user_is_not_changed_except_empty_bale_username(h: Harness) -> None:
    h.register_group()
    staff_id = h.root.add_user(USER["id"], first="محمد", last="کریمی", bale_username=None)
    before = h.root.rows("SELECT FirstName, LastName, Phone, Username, PasswordHash, IsAdmin FROM Users")
    start = len(EXECUTED_SQL)
    await h.text(GROUP, USER, "متن")
    await h.confirm([1])
    assert h.root.rows("SELECT FirstName, LastName, Phone, Username, PasswordHash, IsAdmin FROM Users") == before
    assert h.root.rows("SELECT BaleUsername FROM Users") == [("ali",)]  # only the empty field filled
    assert h.count("Users") == 1
    updates = _user_updates(start)
    assert len(updates) == 1 and "SET BaleUsername" in updates[0]
    assert h.root.rows("SELECT AuthorUserId FROM Posts") == [(staff_id,)]

    # second message: BaleUsername already set → no UPDATE on Users at all
    start = len(EXECUTED_SQL)
    await h.text(GROUP, {**USER, "username": "ali_renamed"}, "باز هم")
    assert _user_updates(start) == []
    assert h.root.rows("SELECT BaleUsername FROM Users") == [("ali",)]


async def test_inactive_or_deleted_user_is_refused(h: Harness) -> None:
    h.register_group()
    h.root.add_user(USER["id"], is_active=0)
    h.root.add_user(OTHER_USER["id"], is_deleted=1)
    await h.text(GROUP, USER, "متن")
    await h.text(GROUP, OTHER_USER, "متن")
    assert len(h.ctx.wizards) == 0
    assert h.count("Posts") == 0
    assert fa.PERSON_INACTIVE in h.texts_to(USER["id"])
    assert fa.PERSON_INACTIVE in h.texts_to(OTHER_USER["id"])


# ─── «📝 توضیحات» ───


async def test_note_reply_is_saved_and_not_archived(h: Harness) -> None:
    h.register_group()
    await h.text(GROUP, USER, "اصل پیام")
    sid = h.sid()
    await h.press(USER, "y", sid)
    labels = h.fake.markup_labels(h.fake.last_markup(USER["id"]))
    assert fa.BTN_NOTE in labels
    prompt = await _note_prompt(h, sid)
    await _reply(h, prompt, "توضیح من")
    assert len(h.ctx.wizards) == 1  # no new wizard for the note itself
    s = h.ctx.wizards.get(sid)
    assert s is not None and s.note == "توضیح من" and s.note_prompt_id is None
    assert (USER["id"], prompt) in h.deleted()  # prompt removed once answered
    labels = h.fake.markup_labels(h.fake.last_markup(USER["id"]))
    assert fa.BTN_NOTE_SET in labels
    await h.press(USER, "t", sid, "1")
    await h.press(USER, "c", sid)
    preview = [m.text for m in h.fake.messages.values() if m.chat_id == USER["id"] and m.reply_markup]
    assert "📝 توضیحات: توضیح من" in (preview[-1] or "")
    await h.press(USER, "f", sid)
    assert h.root.rows("SELECT Body, UserNote FROM Posts") == [("اصل پیام", "توضیح من")]
    assert h.count("Posts") == 1


async def test_note_can_be_edited_and_deleted_in_preview(h: Harness) -> None:
    h.register_group()
    await h.text(GROUP, USER, "متن")
    sid = h.sid()
    await h.press(USER, "y", sid)
    await h.press(USER, "t", sid, "1")
    await h.press(USER, "c", sid)
    await _reply(h, await _note_prompt(h, sid), "نسخه‌ی اول")
    await _reply(h, await _note_prompt(h, sid), "نسخه‌ی دوم")
    assert h.ctx.wizards.get(sid).note == "نسخه‌ی دوم"  # type: ignore[union-attr]
    labels = h.fake.markup_labels(h.fake.last_markup(USER["id"]))
    assert fa.BTN_NOTE_EDIT in labels and fa.BTN_NOTE_DELETE in labels
    await h.press(USER, "nd", sid)
    await h.press(USER, "f", sid)
    assert h.root.rows("SELECT UserNote FROM Posts") == [(None,)]


async def test_note_reply_from_another_user_is_ignored(h: Harness) -> None:
    h.register_group()
    await h.text(GROUP, USER, "متن")
    sid = h.sid()
    await h.press(USER, "y", sid)
    prompt = await _note_prompt(h, sid)
    await _reply(h, prompt, "من کس دیگری هستم", sender=OTHER_USER)
    assert h.ctx.wizards.get(sid).note is None  # type: ignore[union-attr]


async def test_note_over_1000_characters_is_refused(h: Harness) -> None:
    h.register_group()
    await h.text(GROUP, USER, "متن")
    sid = h.sid()
    await h.press(USER, "y", sid)
    prompt = await _note_prompt(h, sid)
    await _reply(h, prompt, "ب" * 1001)
    assert h.ctx.wizards.get(sid).note is None  # type: ignore[union-attr]
    assert fa.NOTE_TOO_LONG in h.texts_to(USER["id"])
    await _reply(h, prompt, "ب" * 1000)  # the same prompt still works
    assert h.ctx.wizards.get(sid).note == "ب" * 1000  # type: ignore[union-attr]


async def test_note_without_hashtag_cannot_be_confirmed(h: Harness) -> None:
    h.register_group()
    await h.text(GROUP, USER, "متن")
    sid = h.sid()
    await h.press(USER, "y", sid)
    await _reply(h, await _note_prompt(h, sid), "فقط توضیح")
    await h.press(USER, "c", sid)  # no hashtag selected
    await h.press(USER, "f", sid)
    assert h.count("Posts") == 0
    assert h.ctx.wizards.get(sid).step == "tags"  # type: ignore[union-attr]


async def test_note_message_in_a_group_is_deleted_when_possible(h: Harness) -> None:
    h.register_group()
    await h.text(GROUP, USER, "متن")
    sid = h.sid()
    await h.press(USER, "y", sid)
    prompt = await _note_prompt(h, sid)
    # A client that lets the user answer the prompt from the group chat.
    h.ctx.wizards.get(sid).note_prompt_id = prompt  # type: ignore[union-attr]
    original = h.ctx.wizards.find_by_note_prompt
    h.ctx.wizards.find_by_note_prompt = lambda chat, mid: original(USER["id"], mid)  # type: ignore[method-assign]
    note_id = await h.send(GROUP, USER, text="توضیح از گروه", reply_to_message={"message_id": prompt, "chat": GROUP, "date": 1})
    assert h.ctx.wizards.get(sid).note == "توضیح از گروه"  # type: ignore[union-attr]
    assert (GROUP_ID, note_id) in h.deleted()
    assert len(h.ctx.wizards) == 1  # not archived, no second wizard


async def test_no_or_timeout_with_note_writes_nothing(h: Harness) -> None:
    h.register_group()
    await h.text(GROUP, USER, "متن")
    sid = h.sid()
    await h.press(USER, "y", sid)
    await _reply(h, await _note_prompt(h, sid), "توضیح")
    await h.press(USER, "x", sid)
    await h.text(GROUP, USER, "دومی")
    sid2 = h.sid()
    await h.press(USER, "y", sid2)
    h.ctx.wizards.get(sid2).created_at -= 31 * 60  # type: ignore[union-attr]
    await expire_and_remind(h.ctx)
    assert h.count("Posts") == 0 and h.count("PostHashtags") == 0 and h.count("PostMedia") == 0


# ─── Source='bale' only ───


def _staff_post(h: Harness, author: int, group_id: int, body: str = "پست نیرو") -> int:
    h.root.execute(
        "INSERT INTO Posts (AuthorUserId, BaleGroupId, BaleMessageId, ContentType, Body, PostedAt, "
        "CreatedAt, OrgUnitId, Source) VALUES (%s, %s, NULL, 1, %s, NOW(), NOW(), 1, 'web')",
        (author, group_id, body),
    )
    post_id = int(h.root.execute("SELECT LAST_INSERT_ID()")[0][0])
    h.root.execute("INSERT INTO PostHashtags (PostId, HashtagId) VALUES (%s, 1)", (post_id,))
    return post_id


async def test_reports_and_excel_only_see_bale_posts(h: Harness) -> None:
    group_id = h.register_group()
    h.fake.file_bytes["d1"] = b"%PDF-1.4"
    await h.send(GROUP, USER, document={"file_id": "d1", "file_name": "a.pdf", "file_size": 8}, caption="سند")
    sid = h.sid()
    await h.press(USER, "y", sid)
    await h.press(USER, "t", sid, "1")
    await _reply(h, await _note_prompt(h, sid), "توضیح اکسل")
    await h.press(USER, "c", sid)
    await h.press(USER, "f", sid)
    await h.settle()
    (user_id,), = h.root.rows("SELECT Id FROM Users")
    _staff_post(h, user_id, group_id)
    assert h.count("Posts") == 2
    async with h.db.tx() as conn:
        report = await build_report(conn)
        data, count = await build_export(conn, PostFilter(), h.settings)
    assert report.total_posts == 1 and dict(report.by_hashtag)["یادگیری"] == 1
    assert count == 1
    sheet = load_workbook(io.BytesIO(data))["آرشیو"]
    headers = [c.value for c in sheet[1]]
    assert "توضیحات" in headers
    assert sheet.cell(row=2, column=headers.index("توضیحات") + 1).value == "توضیح اکسل"
    (media_id,), = h.root.rows("SELECT Id FROM PostMedia")
    link = sheet.cell(row=2, column=headers.index("فایل ۱") + 1)
    assert link.hyperlink.target == build_media_url(media_id, h.settings)
    await h.private(USER, fa.BTN_MY_POSTS)
    assert "📝 توضیح اکسل" in h.texts_to(USER["id"])[-1]


async def test_undo_only_own_bale_post(h: Harness) -> None:
    group_id = h.register_group()
    await h.text(GROUP, USER, "متن")
    await h.confirm([1])
    (user_id,), = h.root.rows("SELECT Id FROM Users")
    staff_post = _staff_post(h, user_id, group_id)  # newer, but not a bale post
    async with h.db.tx() as conn:
        from app.db import repo

        assert await repo.delete_own_post(conn, staff_post, user_id) is None
    (bale_post,), = h.root.rows("SELECT Id FROM Posts WHERE Source = 'bale'")
    await h.private(USER, fa.BTN_UNDO_LAST)
    await h.press(USER, "un", "", str(bale_post))
    assert h.root.rows("SELECT Id, Source FROM Posts") == [(staff_post, "web")]


def test_posted_at_is_tehran_time() -> None:
    from app.timeutil import tehran_from_unix

    assert tehran_from_unix(1_758_880_800) == datetime(2025, 9, 26, 13, 30)
